"""Chuyển file Excel "CÔNG CỤ TÍNH GIÁ KHUYẾN MẠI.xlsx" thành dữ liệu gốc (seed) cho công cụ web.

Cách dùng:
    pip install openpyxl
    python scripts/import_excel.py "<đường dẫn file xlsx>"

Kết quả:
    src/data/seed.js            dữ liệu gốc dạng JS module (admin có thể sửa tiếp trên web)
    tests/fixtures/excel_cases.json  các ca tính sẵn trong sheet 09_ENGINE + 12_CALCULATOR_SALE
                                     để đối chiếu engine JS với Excel
"""
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Thứ tự + mã rút gọn lấy từ công thức SWITCH ở 12_CALCULATOR_SALE!J3 và K3.
# Bảo hành không tính cho Module A, Module B, BT5.5 (12_CALCULATOR_SALE!K15).
COURSE_CODES = [
    ("Module A", "MA", False),
    ("Module B", "MB", False),
    ("BS", "BS", True),
    ("PRE", "PRE", True),
    ("IE4.5", "IE45", True),
    ("IE6.5(M1)", "M1", True),
    ("IE6.5(M2)", "M2", True),
    ("IE7.0", "IE70", True),
    ("BT5.5", "BT55", False),
    ("Ready", "READY", True),
]


def num(v, default=0):
    if v is None or v == "":
        return default
    return float(v) if isinstance(v, float) and not v.is_integer() else int(v)


def rows(ws, ncols, start=2):
    for r in ws.iter_rows(min_row=start, values_only=True):
        r = list(r[:ncols]) + [None] * max(0, ncols - len(r))
        if r[0] in (None, ""):
            continue
        yield r


PROGRAM_NAMES = ["Định kỳ đặc biệt", "Thấp điểm", "Văn phòng mới", "Vũ Đài Trí Tuệ"]


def lkt_base(group):
    """"LKT Thấp điểm" → "Thấp điểm"; "LKT Toàn hệ thống" → "Định kỳ đặc biệt"."""
    rest = group[4:].strip() if group.upper().startswith("LKT ") else group
    return next((p for p in PROGRAM_NAMES if p.lower() == rest.lower()), "Định kỳ đặc biệt")


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    courses = [
        {"label": r[0], "code": r[1], "name": r[2], "mode": r[3], "price": num(r[4]), "voucher": num(r[5])}
        for r in rows(wb["01_DM_KHOA_HOC"], 6)
    ]
    program_vouchers = [
        {"label": r[0], "program": r[1], "code": r[2], "mode": r[3], "voucher": num(r[4])}
        for r in rows(wb["01B_VOUCHER_CTKM"], 5)
    ]
    warranty = [{"name": r[0], "rate": num(r[1])} for r in rows(wb["02_DM_BAO_HANH"], 2)]
    periodic = [
        {"label": r[0], "from": num(r[1]), "to": num(r[2]), "pct": num(r[3]), "cashback": num(r[4])}
        for r in rows(wb["03_CTKM_DINH_KY"], 5)
    ]
    special = [
        {"label": r[0], "method": r[1], "term": num(r[2]), "pct": num(r[3]), "applyType": r[4],
         "minCourses": num(r[5], 1), "note": r[6] or ""}
        for r in rows(wb["04_CTKM_DAC_BIET"], 7)
    ]
    # 04B: Nhãn | Nhóm LKT | Format LKT | Mức cắt (từ 30/10/2026 có thêm cột Nhóm LKT)
    lkt = [{"label": r[0], "group": str(r[1]).strip(), "format": r[2], "cut": num(r[3])}
           for r in rows(wb["04B_CTKM_LKT"], 4)]
    extended = [
        {"label": r[0], "program": r[1], "method": r[2], "term": num(r[3]), "pct": num(r[4]),
         "applyType": r[5], "minCourses": num(r[6], 1), "note": r[7] or ""}
        for r in rows(wb["04C_CTKM_MO_RONG"], 8)
    ]
    citime_routes = [
        {"label": r[0], "group": r[1], "count": num(r[2]), "code": r[3], "name": r[4],
         "status": r[5], "maxTerm": num(r[6])}
        for r in rows(wb["05_CITIME_LO_TRINH"], 7)
    ]
    citime_terms = [num(r[0]) for r in rows(wb["06_DM_KY_HAN_CITIME"], 1)]
    gifts = [
        {"label": r[0], "program": r[1], "milestone": num(r[2]), "method": r[3], "content": r[4], "value": num(r[5])}
        for r in rows(wb["07_QUA_TANG"], 6)
    ]
    routes = []
    for r in rows(wb["08_LO_TRINH"], 20):
        items = [{"code": r[i], "mode": r[i + 1]} for i in range(4, 20, 2) if r[i]]
        routes.append({"label": r[0], "output": str(r[1]) if r[1] is not None else "", "code": r[2],
                       "name": r[3], "items": items})

    sacombank_terms = sorted({s["term"] for s in special if s["method"] == "Sacombank"})
    labels = sorted({c["label"] for c in courses}, key=["IELTS Mentor", "CLASSUP 1-8", "CLASSUP 1-1"].index)
    settings = {
        "campaignName": "Chương trình khuyến mại tháng 10/2026",
        "validFrom": "2026-10-01",
        "validTo": "2026-10-31",
        "labels": [
            {"name": l, "blockedMethods": ["Sacombank", "Citime"] if l == "CLASSUP 1-1" else []}
            for l in labels
        ],
        "programs": [
            {"name": "Định kỳ", "type": "DINH_KY"},
            {"name": "Định kỳ đặc biệt", "type": "DAC_BIET"},
            {"name": "Liên kết trường", "type": "LKT"},
            {"name": "Thấp điểm", "type": "MO_RONG"},
            {"name": "Văn phòng mới", "type": "MO_RONG"},
            {"name": "Vũ Đài Trí Tuệ", "type": "MO_RONG"},
        ],
        "paymentMethods": [
            {"name": "Full", "terms": [0], "giftGroup": "Full", "citime": False, "allowSingle": True},
            {"name": "Giai đoạn", "terms": [0], "giftGroup": "Giai đoạn", "citime": False, "allowSingle": False},
            {"name": "Sacombank", "terms": sacombank_terms, "giftGroup": "Ngân hàng", "citime": False, "allowSingle": False},
            {"name": "Citime", "terms": citime_terms, "giftGroup": "Citime", "citime": True, "allowSingle": False},
        ],
        # Nhóm LKT → chương trình nền (12_CALCULATOR_SALE!K17): % LKT = % chương trình nền − mức cắt
        "lktGroups": [
            {"name": g, "baseProgram": lkt_base(g)} for g in dict.fromkeys(r["group"] for r in lkt)
        ],
        "courseCodes": [
            {"code": c, "short": s, "order": i + 1, "warranty": w} for i, (c, s, w) in enumerate(COURSE_CODES)
        ],
        "note": "Các khoá Module A, Module B, BT5.5, IE7.0 chỉ có hình thức Online. "
                "Các khoá còn lại có cả 2 hình thức Offline/Online.",
    }

    data = {
        "version": 1,
        "settings": settings,
        "courses": courses,
        "programVouchers": program_vouchers,
        "warranty": warranty,
        "periodic": periodic,
        "special": special,
        "lkt": lkt,
        "extended": extended,
        "citimeRoutes": citime_routes,
        "gifts": gifts,
        "routes": routes,
    }
    out = os.path.join(ROOT, "src", "data", "seed.js")
    with open(out, "w", encoding="utf-8") as f:
        f.write("// Sinh tự động bởi scripts/import_excel.py — dữ liệu gốc ban đầu, admin sửa tiếp trên web.\n")
        f.write("export default ")
        json.dump(data, f, ensure_ascii=False, indent=1)
        f.write(";\n")
    print("Đã ghi", out, {k: len(v) for k, v in data.items() if isinstance(v, list)})

    # --- Ca kiểm thử lấy từ giá trị Excel đã tính sẵn ---
    route_by_key = {(r["label"], r["code"]): r for r in routes}
    cases = []
    for r in wb["09_ENGINE"].iter_rows(min_row=2, values_only=True):
        if not r[1] or r[13] != "OK":
            continue
        route = route_by_key.get((r[1], r[4]))
        if not route:
            continue
        method = "Full" if r[7] == "Chung" else r[7]
        cases.append({
            "source": "09_ENGINE", "label": r[1], "program": r[2], "route": r[4], "items": route["items"],
            "method": method, "term": num(r[8]),
            "expect": {"gross": num(r[9]), "voucher": num(r[10]), "pct": num(r[14]), "cashback": num(r[15]),
                       "afterPromo": num(r[16]), "silver": num(r[17]), "gold": num(r[18]),
                       "diamond": num(r[19]), "warrantyBase": num(r[23])},
        })
    ws = wb["12_CALCULATOR_SALE"]
    items = [{"code": ws[f"E{i}"].value, "mode": ws[f"F{i}"].value} for i in range(3, 12) if ws[f"E{i}"].value]
    cases.append({
        "source": "12_CALCULATOR_SALE", "label": ws["B2"].value, "program": ws["B3"].value,
        "method": ws["B4"].value, "term": num(ws["B5"].value), "warranty": ws["B6"].value,
        "lktGroup": ws["B7"].value, "lktFormat": ws["B8"].value, "items": items,
        "expect": {"gross": num(ws["B13"].value), "voucher": num(ws["B14"].value), "pct": num(ws["B15"].value),
                   "cashback": num(ws["B16"].value), "afterPromo": num(ws["B18"].value),
                   "warrantyFee": num(ws["B20"].value), "final": num(ws["B21"].value), "status": ws["B23"].value},
    })
    os.makedirs(os.path.join(ROOT, "tests", "fixtures"), exist_ok=True)
    with open(os.path.join(ROOT, "tests", "fixtures", "excel_cases.json"), "w", encoding="utf-8") as f:
        json.dump(cases, f, ensure_ascii=False)
    print("Đã ghi", len(cases), "ca kiểm thử")


if __name__ == "__main__":
    main(sys.argv[1])
